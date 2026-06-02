# -*- coding: utf-8 -*-
"""
母乳喂养科普 PPT 生成脚本
依赖: pip install python-pptx
运行: python generate_ppt.py
输出: 母乳喂养科普.pptx

设计说明:
- 16:9 宽屏
- 暖色系主题(柔和粉 / 蜜桃 / 青绿),贴合母婴温馨氛围
- 每张幻灯片 = 视频一个分镜:标题 + 要点 + 备注栏(可放解说词)
- 解说词写入"演讲者备注",方便录制配音时使用
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# ---------- 主题配色 ----------
PRIMARY   = RGBColor(0xE2, 0x6D, 0x7E)   # 主色:柔和玫红
SECONDARY = RGBColor(0xF4, 0xA9, 0x88)   # 蜜桃橙
ACCENT    = RGBColor(0x4F, 0xB0, 0xA5)   # 青绿(强调/勾选)
DARK      = RGBColor(0x3A, 0x3A, 0x3A)   # 正文深灰
LIGHT     = RGBColor(0xFF, 0xFF, 0xFF)   # 白
BG_SOFT   = RGBColor(0xFD, 0xF3, 0xF1)   # 浅粉背景
CARD_BG   = RGBColor(0xFF, 0xFB, 0xFA)   # 卡片背景

FONT = "Microsoft YaHei"  # 微软雅黑,缺失时系统会回退

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _txt(frame, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    _solid(shp, color)
    return shp


def add_round_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    _solid(shp, color)
    return shp


def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BG_SOFT)    # 背景底色
    return s


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------- 各类幻灯片模板 ----------
def cover_slide(prs, title, subtitle, footer):
    s = blank_slide(prs)
    # 顶部色带
    add_rect(s, 0, 0, SLIDE_W, Inches(0.35), PRIMARY)
    add_rect(s, 0, SLIDE_H - Inches(0.35), SLIDE_W, Inches(0.35), SECONDARY)
    # 装饰圆点(用圆形)
    from pptx.enum.shapes import MSO_SHAPE
    for (cx, cy, d, col) in [(11.6, 1.0, 1.6, SECONDARY), (12.4, 5.4, 1.0, ACCENT),
                             (1.0, 5.9, 0.8, PRIMARY)]:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d))
        _solid(c, col)
    # 标题
    tb = s.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(10.5), Inches(1.6))
    _txt(tb.text_frame, title, 44, PRIMARY, bold=True)
    # 副标题
    sb = s.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(10.5), Inches(1.0))
    _txt(sb.text_frame, subtitle, 22, DARK)
    # 页脚
    fb = s.shapes.add_textbox(Inches(1.0), Inches(6.55), Inches(11), Inches(0.6))
    _txt(fb.text_frame, footer, 14, RGBColor(0x88, 0x88, 0x88))
    return s


def section_header(slide, index, title):
    """左上角:序号徽标 + 标题"""
    badge = add_round_rect(slide, Inches(0.7), Inches(0.55), Inches(0.95), Inches(0.95), PRIMARY)
    bf = badge.text_frame
    bf.word_wrap = True
    p = bf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = index
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = LIGHT; r.font.name = FONT
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    tb = slide.shapes.add_textbox(Inches(1.85), Inches(0.62), Inches(10.6), Inches(0.95))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _txt(tb.text_frame, title, 30, DARK, bold=True)
    # 标题下划线
    add_rect(slide, Inches(1.9), Inches(1.62), Inches(3.2), Pt(4), SECONDARY)


def bullets_slide(prs, index, title, bullets, notes=""):
    s = blank_slide(prs)
    section_header(s, index, title)
    box = s.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.9))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (lead, desc) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r1 = p.add_run(); r1.text = "● " + lead
        r1.font.size = Pt(22); r1.font.bold = True
        r1.font.color.rgb = PRIMARY; r1.font.name = FONT
        if desc:
            r2 = p.add_run(); r2.text = "  " + desc
            r2.font.size = Pt(18); r2.font.color.rgb = DARK; r2.font.name = FONT
    if notes:
        set_notes(s, notes)
    return s


def cards_slide(prs, index, title, cards, notes=""):
    """2x2 卡片布局,适合'常见问题/误区辟谣'"""
    s = blank_slide(prs)
    section_header(s, index, title)
    positions = [(0.9, 2.05), (7.05, 2.05), (0.9, 4.55), (7.05, 4.55)]
    cw, ch = Inches(5.4), Inches(2.2)
    for (card, (x, y)) in zip(cards, positions):
        head, body = card
        rect = add_round_rect(s, Inches(x), Inches(y), cw, ch, CARD_BG)
        rect.line.color.rgb = SECONDARY
        rect.line.width = Pt(1.5)
        # 标题条
        bar = add_round_rect(s, Inches(x), Inches(y), cw, Inches(0.6), ACCENT)
        bf = bar.text_frame; bf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _txt(bf, head, 17, LIGHT, bold=True, align=PP_ALIGN.CENTER)
        # 正文
        tb = s.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.7), Inches(4.9), Inches(1.4))
        _txt(tb.text_frame, body, 15, DARK)
    if notes:
        set_notes(s, notes)
    return s


def table_slide(prs, index, title, headers, rows, notes=""):
    """表格幻灯片,适合'母乳储存时间'"""
    s = blank_slide(prs)
    section_header(s, index, title)
    nrows, ncols = len(rows) + 1, len(headers)
    gtable = s.shapes.add_table(nrows, ncols, Inches(1.2), Inches(2.2),
                                Inches(10.9), Inches(0.7 * nrows)).table
    for j, h in enumerate(headers):
        cell = gtable.cell(0, j)
        cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        para.runs[0].font.size = Pt(18); para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = LIGHT; para.runs[0].font.name = FONT
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = gtable.cell(i, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if i % 2 else BG_SOFT
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            para.runs[0].font.size = Pt(16)
            para.runs[0].font.color.rgb = DARK; para.runs[0].font.name = FONT
    if notes:
        set_notes(s, notes)
    return s


def closing_slide(prs, title, lines, footer):
    s = blank_slide(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, PRIMARY)
    tb = s.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(1.2))
    _txt(tb.text_frame, title, 40, LIGHT, bold=True, align=PP_ALIGN.CENTER)
    box = s.shapes.add_textbox(Inches(1.5), Inches(3.4), Inches(10.3), Inches(2.4))
    tf = box.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(10)
        r = p.add_run(); r.text = line
        r.font.size = Pt(20); r.font.color.rgb = LIGHT; r.font.name = FONT
    fb = s.shapes.add_textbox(Inches(1.2), Inches(6.6), Inches(10.9), Inches(0.6))
    _txt(fb.text_frame, footer, 14, RGBColor(0xFF, 0xE6, 0xE0), align=PP_ALIGN.CENTER)
    return s


# ======================= 构建演示文稿 =======================
def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 01 封面
    cover_slide(
        prs,
        "母乳喂养",
        "给宝宝最好的第一份礼物  ·  妇产科健康科普",
        "适用对象:孕产妇及家属  |  本课件为健康科普,不替代医生诊疗,具体请遵医嘱",
    )

    # 02 目录/导语
    bullets_slide(prs, "00", "WHO 喂养建议 · 本期内容", [
        ("出生 1 小时内", "尽早开始第一次母乳喂养"),
        ("0–6 个月", "坚持纯母乳喂养,无需额外喂水"),
        ("6 个月后", "添加辅食,母乳可持续到 2 岁及以上"),
        ("本期讲什么", "好处 · 初乳与三早 · 姿势含接 · 按需哺乳 · 常见问题 · 储存 · 误区"),
    ], notes="世界卫生组织和联合国儿童基金会建议:出生 1 小时内开始母乳喂养,前 6 个月纯母乳喂养,之后在添加辅食基础上可持续至 2 岁甚至更久。")

    # 03 对宝宝的好处
    bullets_slide(prs, "01", "母乳对宝宝的好处", [
        ("增强免疫力", "富含抗体,减少腹泻、肺炎、中耳炎"),
        ("好消化好吸收", "呵护娇嫩肠胃,营养量身定制"),
        ("促进大脑发育", "优质脂肪与营养素助力大脑、视力"),
        ("降低远期风险", "肥胖、过敏及某些慢性病风险更低"),
    ], notes="母乳是为宝宝量身定制的天然食物,含丰富抗体,帮助抵御感染;好消化、好吸收;促进大脑视力发育;远期降低肥胖、过敏风险。")

    # 04 对妈妈的好处
    bullets_slide(prs, "02", "母乳对妈妈的好处", [
        ("促进子宫恢复", "吸吮促进宫缩,减少产后出血"),
        ("帮助产后塑形", "消耗孕期储存脂肪"),
        ("降低患癌风险", "降低乳腺癌、卵巢癌风险"),
        ("方便又经济", "温度刚好,随时可喂,省心省钱"),
    ], notes="宝宝吸吮促进子宫收缩、帮助恢复;哺乳消耗脂肪助塑形;长期哺乳降低乳腺癌卵巢癌风险;母乳方便经济。")

    # 05 初乳
    bullets_slide(prs, "03", "初乳:第一剂“天然疫苗”", [
        ("金黄而珍贵", "产后头几天分泌,颜色金黄,量少"),
        ("天然免疫屏障", "富含免疫球蛋白,肠道形成保护"),
        ("切勿丢弃", "别因量少挤掉,让宝宝尽早吃到"),
    ], notes="初乳被称为宝宝第一剂天然疫苗,富含免疫球蛋白,务必让宝宝尽早吃到,不要因量少而丢弃。")

    # 06 三早
    bullets_slide(prs, "04", "“三早”原则:越吸越有", [
        ("早接触", "出生后尽早肌肤接触"),
        ("早吸吮", "出生 1 小时内第一次吸吮"),
        ("早开奶", "吸吮是信号,越早越频繁,奶来得越快"),
    ], notes="提倡早接触、早吸吮、早开奶。宝宝吸吮像信号,刺激越早越频繁,泌乳越快越足。")

    # 07 哺乳姿势
    bullets_slide(prs, "05", "正确的哺乳姿势", [
        ("常用姿势", "摇篮式 · 橄榄球式(环抱式) · 侧卧式"),
        ("三贴原则", "胸贴胸 · 腹贴腹 · 下巴贴乳房"),
        ("一条直线", "宝宝头、颈、身体保持直线"),
        ("放松省力", "用枕头支撑背部和手臂"),
    ], notes="舒适姿势是成功第一步。摇篮式、橄榄球式、侧卧式皆可;做到三贴,头颈身体成直线,用枕头支撑更省力。")

    # 08 含接
    cards_slide(prs, "06", "正确含接(衔乳)要点", [
        ("✔ 正确含接", "嘴张大,含住乳头+大部分乳晕;下唇外翻;下巴贴乳房;脸颊鼓鼓。"),
        ("✘ 错误含接", "只含乳头 → 吃不饱,易致乳头疼痛、皲裂。"),
        ("为什么重要", "含接对不对,直接决定能否吃饱、妈妈会不会痛。"),
        ("含接不当怎么办", "手指轻压宝宝嘴角,让其松开后重新含接。"),
    ], notes="正确含接:嘴张大含住大部分乳晕,下唇外翻,下巴贴乳房。只含乳头会导致吃不饱和乳头皲裂。含接不当用手指轻压嘴角重新含接。")

    # 09 按需哺乳
    bullets_slide(prs, "07", "按需哺乳:读懂饥饿信号", [
        ("按需哺乳", "不必严格掐钟点,每天约 8–12 次"),
        ("早期信号", "转头觅食 · 吸吮小手 · 咂咂嘴"),
        ("哭是晚信号", "最好在宝宝大哭前就开始喂"),
    ], notes="新生儿按需哺乳,每天约 8-12 次。觅食、吸手、咂嘴是早期饥饿信号;哭是很晚的信号,应在大哭前喂。")

    # 10 判断吃饱
    bullets_slide(prs, "08", "怎么判断宝宝吃饱了", [
        ("看尿布", "一周后每天尿湿 ≥ 6 块,大便规律"),
        ("听吞咽", "吃奶时有“咕咚咕咚”吞咽声"),
        ("看状态", "吃完表情满足,安静入睡"),
        ("看体重", "体重稳步增长是最可靠指标"),
    ], notes="判断吃饱:每天尿湿6块以上、有吞咽声、吃后满足入睡、体重稳步增长。这些正常即说明奶量足够。")

    # 11 常见问题
    cards_slide(prs, "09", "常见问题与对策", [
        ("乳头皲裂", "多因含接不当 → 纠正姿势;哺乳后涂少许乳汁晾干。"),
        ("乳房胀奶", "勤喂、排空乳房;冷敷缓解肿胀。"),
        ("警惕乳腺炎", "红、肿、热、痛甚至发烧 → 及时就医,多可继续哺乳。"),
        ("感觉奶不够", "多为暂时性 → 增加亲喂、有效吸吮、多休息、放松心情。"),
    ], notes="皲裂多因含接不当;胀奶要勤喂排空、冷敷;红肿热痛发烧警惕乳腺炎需就医;奶少多为暂时,靠多吸、多休息、放松改善。拿不准请咨询医生。")

    # 12 饮食生活
    bullets_slide(prs, "10", "哺乳期妈妈的饮食与生活", [
        ("均衡多样", "主食 + 优质蛋白 + 蔬果 + 奶类"),
        ("适量喝水", "不必盲目大补、顿顿浓汤油腻"),
        ("忌烟酒", "咖啡浓茶适量;用药前先告知医生在哺乳"),
        ("休息与心情", "充足睡眠 + 好心情同样重要"),
    ], notes="哺乳期均衡多样饮食,适量喝水,不盲目大补;戒烟酒,用药前告知医生;保证睡眠和好心情。")

    # 13 储存表
    table_slide(prs, "11", "母乳的挤出与储存",
        ["存放方式", "建议时间", "提示"],
        [
            ["室温(约 25℃)", "≤ 4 小时", "标注日期,先存先用"],
            ["冷藏室", "3 – 4 天", "置于内壁,避开冰箱门"],
            ["冷冻", "约 6 个月", "解冻后不可再次冷冻"],
            ["解冻方式", "冷藏 / 温水", "切勿微波,以免破坏营养、烫伤"],
        ],
        notes="储存时间:室温≤4小时,冷藏3-4天,冷冻约6个月。标日期、先存先用;温水或冷藏解冻,勿用微波炉;解冻后别再冷冻。")

    # 14 误区辟谣
    cards_slide(prs, "12", "常见误区辟谣", [
        ("✘ 奶太稀没营养", "✔ 前段稀后段浓,前后都是营养,都要吃到。"),
        ("✘ 感冒生气不能喂", "✔ 普通感冒戴口罩、勤洗手多可继续哺乳。"),
        ("✘ 半岁后母乳没营养", "✔ 母乳一直有营养,半岁后需在母乳基础上加辅食。"),
        ("✘ 攒着奶会更多", "✔ 恰恰相反,排空勤吸才能越喂越多。"),
    ], notes="辟谣:奶稀也有营养;普通感冒多可继续喂;半岁后母乳仍有营养需加辅食;攒奶并不会更多,越吸越有。")

    # 15 结尾
    closing_slide(prs, "科学喂养 · 温暖陪伴", [
        "母乳喂养是妈妈和宝宝共同学习、彼此适应的过程,",
        "需要耐心,也需要家人的支持。",
        "遇到困难别独自硬扛,母乳喂养门诊随时为您提供帮助。",
        "愿每一位妈妈,都能享受这段温暖的时光。",
    ], footer="母乳喂养门诊  |  咨询电话:[ ____ ]  |  [科室名称 · 二维码]")

    out = "母乳喂养科普.pptx"
    prs.save(out)
    print(f"已生成: {out}  (共 {len(prs.slides._sldIdLst)} 张幻灯片)")


if __name__ == "__main__":
    build()
